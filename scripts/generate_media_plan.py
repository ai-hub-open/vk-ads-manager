#!/usr/bin/env python3
"""
generate_media_plan.py — собирает media_plan.docx из всех артефактов рабочей папки.

Если python-docx не установлен — фолбек на media_plan.md.

Usage:
    python -m scripts.generate_media_plan --workspace <path>
"""

import argparse
import json
import sys
from pathlib import Path
from datetime import datetime


ARTIFACTS_ORDER = [
    ("01_brief.md", "1. Продуктовый бриф"),
    ("02_visual.md", "2. Визуал"),
    ("03_competitors.md", "3. Конкуренты"),
    ("04_competitor_analysis.md", "4. Анализ конкурентов"),
    ("05_positioning.md", "5. Позиционирование"),
    ("06_personas.md", "6. Buyer Personas"),
    ("_audiences.md", "6.5. Матрица аудиторий"),
    ("07_competitor_ads.md", "7. Реклама конкурентов"),
    ("07a_pixel_setup.md", "7а. Пиксель и события"),
    ("08_patterns.md", "8. Успешные паттерны"),
    ("09_landing.md", "9. Посадочные"),
    ("10_creatives.md", "10. Креативы"),
]


def render_md(workspace: Path, output: Path):
    """Собирает все артефакты в один markdown."""
    lines = [
        "# Медиаплан VK Реклама",
        "",
        f"Дата сборки: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
        "---",
        "",
    ]

    state_path = workspace / "_state.json"
    if state_path.exists():
        with state_path.open(encoding="utf-8") as f:
            state = json.load(f)
        lines.append(f"**Slug:** {state.get('slug', '?')}")
        lines.append(f"**Текущий шаг:** {state.get('current_step', '?')}")
        lines.append("")

    for fname, title in ARTIFACTS_ORDER:
        fpath = workspace / fname
        if fpath.exists():
            lines.append(f"## {title}")
            lines.append("")
            lines.append(fpath.read_text(encoding="utf-8"))
            lines.append("")
            lines.append("---")
            lines.append("")
        else:
            lines.append(f"## {title}")
            lines.append("")
            lines.append("_(не подготовлено)_")
            lines.append("")
            lines.append("---")
            lines.append("")

    output.write_text("\n".join(lines), encoding="utf-8")


def render_docx(workspace: Path, output: Path) -> bool:
    """Собирает docx через python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return False

    doc = Document()

    title = doc.add_heading("Медиаплан VK Реклама", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Дата сборки: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    state_path = workspace / "_state.json"
    if state_path.exists():
        with state_path.open(encoding="utf-8") as f:
            state = json.load(f)
        p = doc.add_paragraph()
        p.add_run(f"Slug: ").bold = True
        p.add_run(state.get("slug", "?"))
        p = doc.add_paragraph()
        p.add_run(f"Текущий шаг: ").bold = True
        p.add_run(str(state.get("current_step", "?")))

    doc.add_paragraph("─" * 50)

    for fname, title_text in ARTIFACTS_ORDER:
        fpath = workspace / fname
        doc.add_heading(title_text, level=1)
        if fpath.exists():
            text = fpath.read_text(encoding="utf-8")
            # Простой рендеринг markdown как plaintext с разделением на параграфы
            for paragraph in text.split("\n\n"):
                if paragraph.strip():
                    p = doc.add_paragraph(paragraph.strip())
                    p.paragraph_format.space_after = Pt(6)
        else:
            p = doc.add_paragraph("(не подготовлено)")
            run = p.runs[0]
            run.italic = True
            run.font.color.rgb = RGBColor(150, 150, 150)

    doc.save(output)
    return True


def main():
    parser = argparse.ArgumentParser(description="Генерирует медиаплан в docx или md")
    parser.add_argument("--workspace", required=True, help="Папка vk-campaign-<slug>/")
    parser.add_argument("--format", choices=["docx", "md", "auto"], default="auto")
    args = parser.parse_args()

    workspace = Path(args.workspace)
    if not workspace.exists():
        print(f"ERROR: workspace {workspace} не существует", file=sys.stderr)
        return 1

    if args.format in ("docx", "auto"):
        docx_path = workspace / "media_plan.docx"
        if render_docx(workspace, docx_path):
            print(f"OK: {docx_path}")
            return 0
        if args.format == "docx":
            print("ERROR: python-docx не установлен", file=sys.stderr)
            return 1
        # Фолбек

    md_path = workspace / "media_plan.md"
    render_md(workspace, md_path)
    print(f"OK (md фолбек): {md_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
