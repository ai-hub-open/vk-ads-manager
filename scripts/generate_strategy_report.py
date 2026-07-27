#!/usr/bin/env python3
"""
generate_strategy_report.py — собирает отчёт-стратегию для команды агентства.

Единый источник правды: структура кампаний, воронки, посадочные и лид-формы
выводятся из creatives.json + audiences.json — тех же данных, что идут в ads.xlsx.
Поэтому стратегия и таблица кампаний не противоречат друг другу.

Выдаёт report/strategy_report.md (+ .docx с --docx) и самодостаточные .md-артефакты
по блокам (_competitors, _semantics, _usp, _landings, _funnels, _campaign_structure,
_forecast, _client_requests, _strategy).

Usage:
    python -m scripts.generate_strategy_report --workspace <path> [--docx]
"""

import argparse
import json
import sys
from pathlib import Path
from datetime import datetime

STANDALONE_MAP = {
    "04_competitor_analysis.md": "_competitors.md",
    "_semantics.md": "_semantics.md",
    "05_positioning.md": "_usp.md",
    "_usp.md": "_usp.md",
    "09_landing.md": "_landings.md",
    "06_personas.md": "_personas.md",
    "_audiences.md": "_audiences.md",
}

FUNNEL_LABELS = {
    "cold": "Холодный трафик", "warm": "Тёплый трафик", "hot": "Горячий трафик",
    "retarget": "Ретаргет", "retargeting": "Ретаргет",
    "leadmagnet": "Лид-магниты", "lead_magnet": "Лид-магниты",
    "subscribers": "Подписчики сообщества", "subs": "Подписчики сообщества",
}


def load_json(path, default):
    if not path.exists():
        return default
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default
    if isinstance(data, dict):
        for k in ("creatives", "audiences", "items"):
            if k in data:
                return data[k]
    return data


def read_text(path):
    return path.read_text(encoding="utf-8") if path.exists() else None


def funnel_label(stage):
    return FUNNEL_LABELS.get((stage or "").lower().strip(), stage or "—")


def esc(value):
    """Экранирует значения для markdown-таблиц: '|' внутри ячейки ломает таблицу."""
    return str(value if value not in (None, "") else "—").replace("|", "·").replace("\n", " ")


def block_funnels(creatives, audiences):
    by_stage = {}
    for c in creatives:
        st = funnel_label(c.get("funnel_stage"))
        by_stage.setdefault(st, {"groups": set(), "creatives": 0, "landings": set()})
        by_stage[st]["groups"].add(c.get("group_name") or c.get("ad_plan_id") or "—")
        by_stage[st]["creatives"] += 1
        is_lf = "lead" in (c.get("format", "") or "").lower()
        by_stage[st]["landings"].add("VK лид-форма" if is_lf else (c.get("url") or "—"))
    lines = ["## Воронки трафика", "",
             "Какие воронки задействованы и чем наполнены. Если воронки нет в таблице — "
             "её в этом запуске не используем (осознанно).", "",
             "| Воронка | Групп | Креативов | Посадочные / лид-формы |",
             "|---|---|---|---|"]
    for st, d in by_stage.items():
        lines.append(f"| {esc(st)} | {len(d['groups'])} | {d['creatives']} | "
                     f"{', '.join(esc(x) for x in sorted(x for x in d['landings'] if x))} |")
    if not by_stage:
        lines.append("| _(нет данных creatives.json)_ | | | |")
    lines += ["",
              "Типовые воронки VK: холодный трафик (охват/осознание) → ретаргет "
              "(догоняем посетителей и зрителей) → лид-магниты (PDF/чек-лист за контакт) "
              "→ подписчики сообщества (прогрев контентом). Каждая — отдельная цель кампании.", ""]
    return "\n".join(lines)


def block_landings(creatives):
    rows = {}
    for c in creatives:
        is_lf = "lead" in (c.get("format", "") or "").lower()
        key = "VK лид-форма" if is_lf else (c.get("url") or "—")
        rows.setdefault(key, {"is_lf": is_lf, "groups": set(), "n": 0})
        rows[key]["groups"].add(c.get("group_name") or "—")
        rows[key]["n"] += 1
    lines = ["## Посадочные и лид-формы", "",
             "| Куда ведём | Тип | Групп | Креативов |", "|---|---|---|---|"]
    for k, d in rows.items():
        typ = "VK лид-форма (без сайта)" if d["is_lf"] else "Лендинг/сайт"
        lines.append(f"| {esc(k)} | {typ} | {len(d['groups'])} | {d['n']} |")
    if not rows:
        lines.append("| _(нет данных)_ | | | |")
    lines += ["",
              "> Для каждой лид-формы VK зафиксировать: какие поля, какой «успешный экран», "
              "куда уходит лид (CRM/почта), отдаётся ли лид-магнит. Для лендинга — пройден ли "
              "`landing-checklist.md` (пиксель, скорость, форма выше скролла, юр.инфо).", ""]
    return "\n".join(lines)


def block_structure(creatives):
    lines = ["## Структура кампаний (единый источник со ads.xlsx)", "",
             "Каждая строка явно связывает кампанию, группу, сегмент, гипотезу, аудиторию, "
             "воронку, посадочную и креатив — ничего не висит в воздухе.", "",
             "| Кампания | Группа | Сегмент | Гипотеза | Аудитория | Воронка | Посадочная | Креатив | Формат |",
             "|---|---|---|---|---|---|---|---|---|"]
    for c in creatives:
        is_lf = "lead" in (c.get("format", "") or "").lower()
        landing = "VK лид-форма" if is_lf else (c.get("url") or "—")
        lines.append("| {camp} | {grp} | {seg} | {hyp} | {aud} | {fun} | {land} | {name} | {fmt} |".format(
            camp=esc(c.get("campaign_name")), grp=esc(c.get("group_name")),
            seg=esc(c.get("segment")), hyp=esc(c.get("hypothesis")),
            aud=esc(c.get("audience_id")), fun=esc(funnel_label(c.get("funnel_stage"))),
            land=esc(landing), name=esc(c.get("name")), fmt=esc(c.get("format") or "universal")))
    if not creatives:
        lines.append("| _(нет creatives.json)_ | | | | | | | | |")
    groups = {c.get("group_name") or c.get("ad_plan_id") for c in creatives}
    lines += ["",
              f"**Групп всего: {len(groups)}.** Дисциплина бюджета: на тесте держим деньги "
              "на немногих сильных гипотезах (≈ ≥ 1000 ₽/группу/день, иначе алгоритм не выходит "
              "из обучения). Если групп много, а бюджет мал — резать число гипотез, а не размазывать.", ""]
    return "\n".join(lines)


def block_forecast(workspace):
    fc = read_text(workspace / "_forecast.md")
    if not fc:
        return ("## Прогноз\n\n_Прогноз не сформирован._ Запусти `scripts/forecast.py` "
                "после получения CPM/CTR из кабинета (MCP) и CR от клиента. "
                "Лиды/CPL без согласования с маркетологом в отчёт не идут.\n")
    note = ("\n> ⚠️ Цифры лидов/CPL действительны только после согласования с маркетологом. "
            "До этого — ориентир, не обязательство.\n" if "требует согласования" in fc else "")
    return "## Прогноз\n\n" + fc + note


def block_client_requests(workspace, creatives, state):
    reqs = []
    if not state.get("pixel_ready"):
        reqs.append("Подтверждение, что **пиксель VK Ads установлен** и события ловятся "
                    "(без него нет ретаргета, lookalike по посетителям и оптимизации на конверсии).")
    reqs.append("**Реальный CR** (клик→заявка) лендинга или лид-формы из их аналитики — "
                "без него прогноз лидов/CPL недостоверен.")
    reqs.append("**Конверсия заявка→продажа** и средний CPL/CAC из прошлого опыта (для reality-check).")
    vis = read_text(workspace / "02_visual.md")
    if not vis or "не хватает" in (vis or "").lower():
        reqs.append("**Реальные материалы**: логотип в высоком разрешении, фото/скриншоты продукта "
                    "(UI-моки AI делает плохо — нужны живые скрины ≥ 2160×2160), видео/отзывы клиентов.")
    reqs.append("**База для custom-аудиторий**: CSV клиентов (phone/email, ≥ 1000 после матчинга) "
                "и/или доступ к сообществу — для lookalike и ретаргета.")
    reqs.append("**Юр.реквизиты для ОРД**: ИП/ООО + ИНН (поле «О компании»), кто рекламодатель.")
    reqs.append("**Доступы**: токен VK Ads API или роль в кабинете; доступ к лендингу для пикселя/UTM.")
    reqs.append("**Пропускная способность отдела продаж** — сколько лидов в день реально обработают.")
    lines = ["## Что запросить у клиента", "",
             "Собрать до запуска — без этих данных стратегия и прогноз держатся на предположениях.", ""]
    lines += [f"{i}. {r}" for i, r in enumerate(reqs, 1)]
    lines.append("")
    return "\n".join(lines)


def build_report(workspace, creatives, audiences, state):
    slug = state.get("slug", workspace.name)
    parts = [f"# Стратегия VK Реклама — {slug}", "",
             f"Дата: {datetime.now().strftime('%Y-%m-%d %H:%M')}  ·  Документ для команды агентства.", ""]
    pos = read_text(workspace / "05_positioning.md")
    if pos:
        parts += ["## Краткое резюме (позиционирование и УТП)", "", pos.strip(), ""]
    parts += [block_funnels(creatives, audiences),
              block_landings(creatives),
              block_structure(creatives),
              block_forecast(workspace),
              block_client_requests(workspace, creatives, state),
              "## Замечания и статус согласований", "",
              "- [ ] Список конкурентов согласован (вкл. заменителей «снизу»)",
              "- [ ] УТП и боли проработаны ПО СЕГМЕНТАМ (не смешаны)",
              "- [ ] Прогноз лидов/CPL согласован с маркетологом",
              "- [ ] Тексты прошли проверку: нет имён конкурентов, нет «лучше чем X», нет «ни о чём»",
              "- [ ] Структура (кампании/группы/посадочные/креативы) совпадает с ads.xlsx", ""]
    return "\n".join(parts)


def write_standalones(workspace, report_dir, creatives, audiences, state):
    written = []
    for src, dst in STANDALONE_MAP.items():
        txt = read_text(workspace / src)
        if txt and not (report_dir / dst).exists():
            (report_dir / dst).write_text(txt, encoding="utf-8")
            written.append(dst)
    gen = {
        "_funnels.md": block_funnels(creatives, audiences),
        "_campaign_structure.md": block_structure(creatives),
        "_landings.md": block_landings(creatives),
        "_client_requests.md": block_client_requests(workspace, creatives, state),
        "_strategy.md": build_report(workspace, creatives, audiences, state),
    }
    fc = read_text(workspace / "_forecast.md")
    if fc:
        gen["_forecast.md"] = fc
    for name, content in gen.items():
        (report_dir / name).write_text(content, encoding="utf-8")
        written.append(name)
    return written


def to_docx(md_text, output):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return False
    doc = Document()
    for raw in md_text.split("\n"):
        line = raw.rstrip()
        if line.startswith("# "):
            h = doc.add_heading(line[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("|"):
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(0)
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(output)
    return True


def main():
    ap = argparse.ArgumentParser(description="Отчёт-стратегия VK для команды агентства")
    ap.add_argument("--workspace", required=True)
    ap.add_argument("--docx", action="store_true")
    args = ap.parse_args()
    ws = Path(args.workspace)
    if not ws.exists():
        print(f"ERROR: workspace {ws} не существует", file=sys.stderr)
        return 1
    creatives = load_json(ws / "creatives.json", [])
    audiences = load_json(ws / "audiences.json", [])
    state = {}
    if (ws / "_state.json").exists():
        try:
            state = json.loads((ws / "_state.json").read_text(encoding="utf-8"))
        except Exception:
            state = {}
    report_dir = ws / "report"
    report_dir.mkdir(exist_ok=True)
    report_md = build_report(ws, creatives, audiences, state)
    (report_dir / "strategy_report.md").write_text(report_md, encoding="utf-8")
    print(f"OK: {report_dir / 'strategy_report.md'}")
    standalones = write_standalones(ws, report_dir, creatives, audiences, state)
    print(f"Standalone-артефакты: {', '.join(standalones)}")
    if args.docx:
        if to_docx(report_md, report_dir / "strategy_report.docx"):
            print(f"OK: {report_dir / 'strategy_report.docx'}")
        else:
            print("python-docx не установлен — docx пропущен (есть .md)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
