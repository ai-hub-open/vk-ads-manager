#!/usr/bin/env python3
"""
generate_ads_xlsx.py — превращает creatives.json в xlsx (или CSV) для VK Реклама.

Валидирует лимиты VK по форматам и выводит warnings.txt со списком проблем
(превышение лимитов, потенциальные триггеры модерации, юр.чистоту).

Usage:
    python -m scripts.generate_ads_xlsx --workspace <path>
"""

import argparse
import json
import sys
import re
from pathlib import Path

# Лимиты VK Реклама (на 2026 май)
LIMITS = {
    "universal": {"title": 40, "description": 16384, "description_preview": 220,
                  "button_text": 30, "company_info": 115},
    "carousel_card": {"title": 40, "description_long": 2000, "description_short": 90,
                      "description_app": 220},
    "video": {"title": 40, "description": 2000, "duration_sec": 60},
    "lead_form": {"title": 40, "description": 200},
    "collage": {"title": 40, "description_long": 2000, "description_short": 90},
}

MODERATION_TRIGGERS = ["100%", "гарантированно", "гарантия результата", "лучший",
                       "единственный", "самый", "номер 1", "№1", "только сегодня",
                       "осталось n мест"]

# Сравнительные утверждения — риск по ст. 5 ФЗ «О рекламе»
COMPARISON_TRIGGERS = ["лучше чем", "лучше, чем", "дешевле всех", "дешевле чем",
                       "лидер рынка", "превосходит", "круче чем", "никто не делает",
                       "нет аналогов", "вне конкуренции"]

# Шаблонные «ни о чём» фразы
VAGUE_PHRASES = ["качество и индивидуальный подход", "индивидуальный подход",
                 "лучшее решение", "широкий спектр", "достичь целей", "для вашего бизнеса",
                 "высокое качество", "профессиональный подход", "выгодные условия",
                 "надёжный партнёр", "надежный партнер"]


def check_caps(text):
    return len(re.findall(r"\b[А-ЯA-Z]{3,}\b", text or "")) > 0


def check_exclamations(text):
    return "!!" in (text or "")


def check_triggers(text):
    t = (text or "").lower()
    return [x for x in MODERATION_TRIGGERS if x in t]


def check_comparisons(text):
    t = (text or "").lower()
    return [x for x in COMPARISON_TRIGGERS if x in t]


def check_competitor_names(text, names):
    t = (text or "").lower()
    return [n for n in (names or []) if n and n.lower() in t]


def check_vague(title, desc):
    warns = []
    full = f"{title} {desc}".lower()
    found = [p for p in VAGUE_PHRASES if p in full]
    if found:
        warns.append(f"шаблонные фразы «ни о чём»: {', '.join(found)}")
    preview = (desc or title or "")[:220]
    has_signal = bool(re.search(r"\d", preview) or "%" in preview or "₽" in preview
                      or "руб" in preview.lower() or "?" in preview)
    if not has_signal:
        warns.append("нет конкретики в первых 220 знаках (ни числа, ни вопроса, ни оффера) — возможно «ни о чём»")
    return warns


def validate_creative(creative, competitor_names=None):
    warnings = []
    fmt = creative.get("format", "universal")
    name = creative.get("name", "<без имени>")
    fmt_limits = LIMITS.get(fmt, LIMITS["universal"])
    title = creative.get("title", "")
    desc = creative.get("description", "")
    button = creative.get("button_text", "")
    company = creative.get("company_info", "")

    if title and len(title) > fmt_limits.get("title", 40):
        warnings.append(f"[{name}] {fmt}: title {len(title)} > лимит {fmt_limits['title']}")
    desc_limit = fmt_limits.get("description") or fmt_limits.get("description_long")
    if desc and desc_limit and len(desc) > desc_limit:
        warnings.append(f"[{name}] {fmt}: description {len(desc)} > лимит {desc_limit}")
    if button and len(button) > fmt_limits.get("button_text", 30):
        warnings.append(f"[{name}] {fmt}: button_text {len(button)} > 30")
    if company and len(company) > 115:
        warnings.append(f"[{name}] {fmt}: company_info {len(company)} > 115")

    full_text = " ".join([title, desc])
    if check_caps(full_text):
        warnings.append(f"[{name}] {fmt}: КАПС в тексте (>=3 заглавных букв подряд)")
    if check_exclamations(full_text):
        warnings.append(f"[{name}] {fmt}: двойной восклицательный знак")
    triggers = check_triggers(full_text)
    if triggers:
        warnings.append(f"[{name}] {fmt}: триггеры модерации: {', '.join(triggers)}")

    comparisons = check_comparisons(full_text)
    if comparisons:
        warnings.append(f"[{name}] {fmt}: сравнительные утверждения (ст. 5 ФЗ о рекламе): {', '.join(comparisons)}")
    comp_names = check_competitor_names(full_text, competitor_names)
    if comp_names:
        warnings.append(f"[{name}] {fmt}: имя конкурента в тексте (только в таргете/ключах!): {', '.join(comp_names)}")
    for w in check_vague(title, desc):
        warnings.append(f"[{name}] {fmt}: {w}")

    if desc and fmt == "universal" and len(desc) > 220:
        preview = desc[:220]
        if not any(p in preview for p in [".", "!", "?"]):
            warnings.append(f"[{name}] {fmt}: первые 220 знаков описания без точки/завершения — обрыв")

    if not creative.get("ad_plan_id"):
        warnings.append(f"[{name}] {fmt}: не указан ad_plan_id (не привязан к группе)")
    if not creative.get("url") and fmt != "lead_form":
        warnings.append(f"[{name}] {fmt}: не указан url")
    if not creative.get("call_to_action"):
        warnings.append(f"[{name}] {fmt}: не указана CTA-кнопка")
    if not company:
        warnings.append(f"[{name}] {fmt}: пустое company_info (юр.инфо для ОРД)")
    return warnings


CREATIVE_HEADERS = ["campaign_name", "group_name", "segment", "hypothesis", "audience_id",
                    "funnel_stage", "landing_or_leadform", "name", "ad_plan_id", "format",
                    "title", "title_len", "description", "desc_len", "button_text",
                    "call_to_action", "url", "company_info", "image_or_video", "notes"]


def creative_row(c):
    is_lf = "lead" in (c.get("format", "") or "").lower()
    landing = "VK лид-форма" if is_lf else c.get("url", "")
    return [c.get("campaign_name", ""), c.get("group_name", ""), c.get("segment", ""),
            c.get("hypothesis", ""), c.get("audience_id", ""), c.get("funnel_stage", ""),
            landing, c.get("name", ""), c.get("ad_plan_id", ""), c.get("format", "universal"),
            c.get("title", ""), len(c.get("title", "")), c.get("description", "")[:500],
            len(c.get("description", "")), c.get("button_text", ""), c.get("call_to_action", ""),
            c.get("url", ""), c.get("company_info", ""), c.get("image_or_video", ""),
            c.get("notes", "")]


def to_xlsx(creatives, audiences, output_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Креативы"
    ws1.append(CREATIVE_HEADERS)
    for c in creatives:
        ws1.append(creative_row(c))
    fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for col in range(1, len(CREATIVE_HEADERS) + 1):
        ws1.cell(row=1, column=col).font = Font(bold=True)
        ws1.cell(row=1, column=col).fill = fill
    widths = [26, 26, 18, 24, 14, 14, 22, 25, 12, 12, 45, 10, 60, 10, 20, 18, 40, 30, 20, 30]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet(title="Аудитории")
    aud_headers = ["id", "name", "persona", "type", "source", "size_estimate", "purpose", "config_json"]
    ws2.append(aud_headers)
    for a in audiences:
        ws2.append([a.get("id", ""), a.get("name", ""), a.get("persona", ""), a.get("type", ""),
                    a.get("source", ""), a.get("size_estimate", ""), a.get("purpose", ""),
                    json.dumps(a.get("config", {}), ensure_ascii=False)])
    for col in range(1, len(aud_headers) + 1):
        ws2.cell(row=1, column=col).font = Font(bold=True)
        ws2.cell(row=1, column=col).fill = fill
    for i, w in enumerate([8, 30, 8, 15, 35, 15, 25, 60], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    wb.save(output_path)
    return True


def to_csv(creatives, audiences, output_dir):
    import csv
    ads_path = output_dir / "ads.csv"
    with ads_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([h for h in CREATIVE_HEADERS if h not in ("title_len", "desc_len")])
        for c in creatives:
            row = creative_row(c)
            row = [v for i, v in enumerate(row) if CREATIVE_HEADERS[i] not in ("title_len", "desc_len")]
            w.writerow(row)
    aud_path = output_dir / "audiences.csv"
    with aud_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "name", "persona", "type", "source", "size_estimate", "purpose"])
        for a in audiences:
            w.writerow([a.get("id", ""), a.get("name", ""), a.get("persona", ""), a.get("type", ""),
                        a.get("source", ""), a.get("size_estimate", ""), a.get("purpose", "")])
    return ads_path, aud_path


def main():
    ap = argparse.ArgumentParser(description="Генерирует xlsx/CSV с креативами и аудиториями")
    ap.add_argument("--workspace", required=True)
    args = ap.parse_args()
    ws = Path(args.workspace)
    if not ws.exists():
        print(f"ERROR: workspace {ws} не существует", file=sys.stderr)
        return 1
    cpath = ws / "creatives.json"
    if not cpath.exists():
        print(f"ERROR: creatives.json не найден в {ws}", file=sys.stderr)
        return 1
    creatives = json.loads(cpath.read_text(encoding="utf-8"))
    if isinstance(creatives, dict):
        creatives = creatives.get("creatives", [])
    audiences = []
    apath = ws / "audiences.json"
    if apath.exists():
        audiences = json.loads(apath.read_text(encoding="utf-8"))
        if isinstance(audiences, dict):
            audiences = audiences.get("audiences", [])

    competitor_names = []
    bpath = ws / "brand.json"
    if bpath.exists():
        try:
            competitor_names = json.loads(bpath.read_text(encoding="utf-8")).get("competitor_names", []) or []
        except Exception:
            pass

    all_warnings = []
    for c in creatives:
        all_warnings.extend(validate_creative(c, competitor_names))

    wpath = ws / "warnings.txt"
    with wpath.open("w", encoding="utf-8") as f:
        if all_warnings:
            f.write(f"# {len(all_warnings)} предупреждений\n\n")
            for w in all_warnings:
                f.write(f"- {w}\n")
        else:
            f.write("Предупреждений нет — все креативы валидны.\n")

    xlsx_path = ws / "ads.xlsx"
    if to_xlsx(creatives, audiences, xlsx_path):
        print(f"OK: {xlsx_path}")
    else:
        ads_csv, aud_csv = to_csv(creatives, audiences, ws)
        print(f"openpyxl не установлен — фолбек на CSV: {ads_csv}, {aud_csv}")
    print(f"Креативов: {len(creatives)}, аудиторий: {len(audiences)}")
    print(f"Предупреждений: {len(all_warnings)} (см. {wpath})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
